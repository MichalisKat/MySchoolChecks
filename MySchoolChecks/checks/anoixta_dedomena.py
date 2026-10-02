"""
checks/anoixta_dedomena.py
══════════════════════════
Β5. Ανοιχτά Δεδομένα — Στοιχεία δημόσιων σχολικών μονάδων ανά Δήμο.

Το UI είναι το AnoixtaDialog στο main.py (ίδια λογική με το Β2 MonadaDialog:
Λήψη 2.2/2.4/2.5 → επιλογή Δήμου → Excel + προαιρετικό email στον Δήμο).
Το run() παρακάτω μένει ως εναλλακτική εκτέλεση (όλοι οι Δήμοι).

Πηγές (MySchool → Στατιστικά):
  2.2  Εκτεταμένα Στοιχεία Σχολικών Μονάδων  → Δήμος, Κωδικός, Ονομασία,
                                                Ενεργοί Μαθητές (ΜΑΘ. ΔΥΝΑΜΙΚΟ),
                                                Είδος, Αναστολή
  2.4  Χώροι σχολικών μονάδων (ανά αίθουσα)   → ΑΙΘΟΥΣΕΣ ΔΙΔΑΣΚΑΛΙΑΣ (πλήθη)
  2.5  Κτιριακά στοιχεία (ανά κτίριο)         → ΚΤΙΡΙΑΚΑ (Ναι/Όχι, έτος)

Αποτέλεσμα: 1 Excel με ένα φύλλο ανά Δήμο (+ φύλλο «ΟΛΟΙ ΟΙ ΔΗΜΟΙ»), στη
μορφή του πίνακα ανοιχτών δεδομένων της Δ/νσης Δ.Ε. + ΜΙΑ ΓΡΑΜΜΗ ΑΝΑ ΚΤΙΡΙΟ:
νέα στήλη D «ΚΤΙΡΙΟ», οι υπόλοιπες μετατοπίζονται κατά 1 (A..AH). Α/Α, ΚΩΔ.,
ΣΧΟΛΕΙΟ, ΜΑΘ. ΔΥΝΑΜΙΚΟ ενώνονται κάθετα (merge) για τα κτίρια ενός σχολείου.

Κανόνες:
  • Εξαιρούνται Ιδιωτικά σχολεία και σχολεία με Αναστολή = True/NAI.
  • ΑΙΘΟΥΣΕΣ: μετρώνται οι χώροι με «Χρήση αίθουσας» = Αίθουσα Διδασκαλίας.
    Συνολικός αριθμός = όλες· Ενεργών = Ενεργός Χώρος NAI· Λυόμενων = Λυόμενο
    NAI· Ηλεκτρική/Δίκτυο/WiFi/Διαδραστικός = NAI (σε όλες τις αίθουσες).
  • ΚΤΙΡΙΑ: μία γραμμή για κάθε κτίριο του 2.5 με τα δικά του κτιριακά.
    Οι αίθουσες του 2.4 μοιράζονται στα κτίρια με βάση τη στήλη «Κτίριο»
    (= «Ονομασία Κτιρίου» του 2.5). Αίθουσες με κτίριο που δεν βρίσκεται
    στο 2.5 → στο 1ο κτίριο του σχολείου (+ σημείωση στο φύλλο ΕΛΛΕΙΨΕΙΣ).
  • ΣΥΝΟΛΑ: μόνο μαθητές/αίθουσες (άθροισμα)· στα κτιριακά δεν μπαίνει σύνολο.
  • Σχολείο που λείπει από το 2.4 / 2.5 → κενά κελιά στις αντίστοιχες στήλες.
"""

import csv, io, os, re, zipfile
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter

# ── Μεταδεδομένα ────────────────────────────────────────────────────────────
CHECK_TITLE       = 'Ανοιχτά Δεδομένα'
CHECK_DESCRIPTION = 'Στοιχεία δημόσιων σχολικών μονάδων ανά Δήμο (μαθητές, αίθουσες, κτιριακά)'
RESULTS_FOLDER    = 'anoixta_dedomena'
HAS_EMAIL         = False
CUSTOM_RUN        = True
NO_SEND_TAB       = True
REQUIRED_REPORTS  = [
    '2.2 — Εκτεταμένα Στοιχεία Σχολ. Μονάδων',
    '2.4 — Χώροι σχολικών μονάδων',
    '2.5 — Κτιριακά στοιχεία',
]

YES = 'Ναι'
NO  = 'Όχι'

# Στήλες ΚΤΙΡΙΑΚΑ στο 2.5 (θέσεις στηλών, 0-based) → στήλη excel εξόδου
# 2.5: 0 Κωδικός | 1 Σχολική Μονάδα | 2 Ονομασία Κτιρίου | 3 Συγκρότημα |
#      4 Έτος κατασκευής | 5 Μισθωμένο | 6..25 (βλ. παρακάτω)
COLS_25_YN = list(range(5, 26))          # Μισθωμένο .. Μόνωση (21 στήλες: M..AG)
COL_25_YEAR = 4
EXPECTED_25 = {6: 'Οικοδομικά Προβλήματα', 7: 'Ύδρευσης', 25: 'Μόνωση'}

# ── Μορφοποίηση (ίδια με το πρότυπο ανοιχτών δεδομένων) ─────────────────────
def _tc(theme, tint=0.0):
    return PatternFill('solid', fgColor=Color(theme=theme, tint=tint))

F_GREY   = _tc(0, -0.1499984740745262)   # γκρι (Α/Α .. ΜΑΘ. ΔΥΝΑΜΙΚΟ, ΚΤΙΡΙΑΚΑ)
F_DGREY  = _tc(0, -0.4999847407452621)   # σκούρο γκρι (κενά group κελιά γραμμής 4)
F_ORANGE = _tc(5, 0.5999938962981048)    # ΑΙΘΟΥΣΕΣ ΔΙΔΑΣΚΑΛΙΑΣ
F_BLUE   = _tc(8, 0.7999816888943144)    # ΠΑΛΑΙΟΤΗΤΑ ΔΙΚΤΥΩΝ
F_GREEN  = _tc(9, 0.5999938962981048)    # ΚΕΝΤΡΙΚΗ ΘΕΡΜΑΝΣΗ
F_YEL6   = _tc(7, 0.5999938962981048)    # ΑΝΤΙΚΕΡΑΥΝΙΚΗ
F_YEL8   = _tc(7, 0.7999816888943144)    # ΣΥΣΤΗΜΑ ΠΥΡΟΣΒΕΣΗΣ
F_NAVY   = _tc(3, 0.5999938962981048)    # ΕΞΥΠΗΡΕΤΗΣΗ ΑΜΕΑ
F_ORG4   = _tc(5, 0.3999755851924192)    # ΕΝΕΡΓΕΙΑΚΗ ΑΝΑΒΑΘΜΙΣΗ

_THIN   = Side(style='thin')
BORDER  = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
CTR     = Alignment(horizontal='center', vertical='center', wrap_text=True)
VERT    = Alignment(horizontal='center', vertical='center', wrap_text=True, textRotation=90)

# Γραμμή 5: (στήλη, τίτλος, fill)
HDR5 = [
    ('E', 'ΣΥΝΟΛΙΚΟΣ ΑΡΙΘΜΟΣ', F_ORANGE), ('F', 'ΑΡΙΘΜΟΣ ΕΝΕΡΓΩΝ', F_ORANGE),
    ('G', 'ΑΡΙΘΜΟΣ ΛΥΟΜΕΝΩΝ', F_ORANGE), ('H', 'ΜΕ ΗΛΕΚΤΡΙΚΗ ΕΓΚΑΤΑΣΤΑΣΗ', F_ORANGE),
    ('I', 'ΜΕ ΔΙΚΤΥΟ', F_ORANGE), ('J', 'ΜΕ WiFi', F_ORANGE),
    ('K', 'ΜΕ ΔΙΑΔΡΑΣΤΙΚΟ ΠΙΝΑΚΑ', F_ORANGE),
    ('L', 'ΕΤΟΣ ΚΑΤΑΣΚΕΥΗΣ', F_GREY), ('M', 'ΜΙΣΘΩΜΕΝΟ', F_GREY),
    ('N', 'ΟΙΚΟΔΟΜΙΚΑ ΠΡΟΒΛΗΜΑΤΑ', F_GREY),
    ('O', 'ΥΔΡΕΥΣΗΣ', F_BLUE), ('P', 'ΘΕΡΜΑΝΣΗΣ', F_BLUE), ('Q', 'ΗΛΕΚΤΡΙΣΜΟΥ', F_BLUE),
    ('R', 'ΤΗΛΕΦΩΝΙΚΟ-ΙΝΤΕΡΝΕΤ', F_BLUE),
    ('S', 'ΠΕΤΡΕΛΑΙΟ', F_GREEN), ('T', 'ΦΥΣΙΚΟ ΑΕΡΙΟ', F_GREEN),
    ('U', 'ΑΙΘΟΣΕΣ ΔΙΔΑΣΚΑΛΙΑΣ ΜΕ Α/C', F_GREEN),
    ('V', 'ΔΙΑΘΕΤΕΙ ΑΝΤΙΚΕΡΑΥΝΙΚΗ ΠΡΟΣΤΑΣΙΑ', F_YEL6),
    ('W', 'ΠΥΡΟΣΒΕΣΤΗΡΕΣ', F_YEL8),
    ('X', 'ΚΕΝΤΡΙΚΟ ΣΥΣΤΗΜΑ ΠΥΡΟΣΒΕΣΗΣ ΣΕ ΟΛΟ ΤΟ ΚΤΙΡΙΟ', F_YEL8),
    ('Y', 'ΣΥΣΤΗΜΑ ΠΥΡΟΣΒΕΣΗΣ ΣΕ ΜΕΡΟΣ ΤΟΥ ΚΤΙΡΙΟΥ', F_YEL8),
    ('Z', 'ΕΞΥΠΗΡΕΤΗΣΗ ΑΜΕΑ', F_NAVY), ('AA', 'ΡΑΜΠΑ', F_NAVY),
    ('AB', 'ΑΝΕΛΚΥΣΤΗΡΑΣ', F_NAVY), ('AC', 'ΤΟΥΑΛΕΤΑ', F_NAVY), ('AD', 'SOFT ROOM', F_NAVY),
    ('AE', 'ΕΝΕΡΓΕΙΑΚΗ ΑΝΑΒΑΘΜΙΣΗ ΤΗΝ ΤΕΛΕΥΤΑΙΑ 10ΕΤΙΑ', F_ORG4),
    ('AF', 'ΚΟΥΦΩΜΑΤΑ', F_ORG4), ('AG', 'ΜΟΝΩΣΗ', F_ORG4),
]
FILL_BY_COL = {c: f for c, _, f in HDR5}
# Γραμμή 4: ομάδες (merge, τίτλος, fill)
HDR4 = [
    ('L4:N4', None, F_DGREY),
    ('O4:R4', 'ΠΑΛΑΙΟΤΗΤΑ ΔΙΚΤΥΩΝ', F_BLUE),
    ('S4:U4', 'ΚΕΝΤΡΙΚΗ ΘΕΡΜΑΝΣΗ', F_GREEN),
    ('V4:V4', None, F_DGREY),
    ('W4:Y4', 'ΣΥΣΤΗΜΑ ΠΥΡΟΣΒΕΣΗΣ', F_YEL8),
    ('Z4:AD4', 'ΕΞΥΠΗΡΕΤΗΣΗ ΑΜΕΑ', F_NAVY),
    ('AE4:AG4', 'ΕΝΕΡΓEΙΑΚΗ ΑΝΑΒΑΘΜΙΣΗ', F_ORG4),
]
COL_WIDTHS = {'A': 4.5, 'B': 9, 'C': 34, 'D': 11}   # σε στήλες ΠΡΟΤΥΠΟΥ
NARROW = 5.0
LAST_COL = 'AH'          # στήλη εξόδου (πρότυπο AG + 1 για «ΚΤΙΡΙΟ»)
BLD_COL  = 'D'           # νέα στήλη ΚΤΙΡΙΟ
BLD_WIDTH = 24


def oc(tpl_letter):
    """Στήλη ΠΡΟΤΥΠΟΥ (A..AG) → στήλη ΕΞΟΔΟΥ: A..C ίδιες, από D και μετά +1."""
    from openpyxl.utils import column_index_from_string as _ci
    i = _ci(tpl_letter)
    return get_column_letter(i if i <= 3 else i + 1)


def _shift_rng(rng):
    """'L4:N4' (πρότυπο) → 'M4:O4' (έξοδος)."""
    out = []
    for part in rng.split(':'):
        col = re.sub(r'\d', '', part)
        row = re.sub(r'\D', '', part)
        out.append(f'{oc(col)}{row}')
    return ':'.join(out)


# ═══════════════════════════════════════════════════════════════════
# ΑΝΑΓΝΩΣΗ ΑΡΧΕΙΩΝ
# ═══════════════════════════════════════════════════════════════════

def _clean_code(v):
    """'="9190637"' / 9190637 / '9190637.0' → '9190637'."""
    if v is None:
        return ''
    s = str(v).replace('=', '').replace('"', '').strip()
    if s.lower() == 'nan':
        return ''
    if re.fullmatch(r'\d+\.0', s):
        s = s[:-2]
    return s


def _decode(raw):
    """UTF-8 αν γίνεται· αλλιώς ISO-8859-7 ή cp1253. Τα CSV του MySchool είναι
    ISO-8859-7: με cp1253 το «Ά» (0xB6) γίνεται «¶» → επιλέγεται η κωδικοποίηση
    με τα λιγότερα «¶»."""
    try:
        return raw.decode('utf-8-sig')
    except UnicodeDecodeError:
        pass
    cands = []
    for enc in ('iso-8859-7', 'cp1253'):
        txt = raw.decode(enc, errors='replace')
        cands.append((txt.count('¶') + txt.count('\ufffd'), txt))
    return min(cands, key=lambda t: t[0])[1]


def _read_rows(path):
    """Επιστρέφει λίστα γραμμών (list of lists) από .xlsx / .xls / .csv / .zip."""
    low = path.lower()
    if low.endswith('.zip'):
        with zipfile.ZipFile(path) as z:
            name = next((n for n in z.namelist()
                         if n.lower().endswith(('.csv', '.xlsx', '.xls'))), None)
            if name is None:
                raise ValueError(f'Το zip {os.path.basename(path)} δεν περιέχει csv/xlsx.')
            data = z.read(name)
        if name.lower().endswith('.csv'):
            return _csv_rows(_decode(data))
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=os.path.splitext(name)[1], delete=False) as t:
            t.write(data)
            tmp = t.name
        try:
            return _read_rows(tmp)
        finally:
            os.remove(tmp)
    if low.endswith('.csv'):
        with open(path, 'rb') as f:
            return _csv_rows(_decode(f.read()))
    if low.endswith('.xls'):
        df = pd.read_excel(path, header=None, dtype=object)
        return [[None if pd.isna(v) else v for v in r] for r in df.itertuples(index=False)]
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _csv_rows(text):
    sample = text[:5000]
    sep = ';' if sample.count(';') >= sample.count(',') else ','
    return [r for r in csv.reader(io.StringIO(text), delimiter=sep)]


def _s(v):
    return '' if v is None else str(v).strip()


def _find_header_row(rows, needle, max_scan=6):
    """Τελευταία γραμμή (από τις πρώτες max_scan) που περιέχει το needle."""
    found = None
    for i, r in enumerate(rows[:max_scan]):
        if any(_s(v) == needle for v in r):
            found = i
    return found


def load_22(path):
    """2.2 → DataFrame [code, name, dimos, eidos, students]
    (χωρίς Ιδιωτικά και σχολεία σε Αναστολή)."""
    rows = _read_rows(path)
    hi = _find_header_row(rows, 'Κωδ. ΥΠΠΘ')
    if hi is None:
        raise ValueError('Το 2.2 δεν έχει στήλη «Κωδ. ΥΠΠΘ».')
    hdr = [_s(v) for v in rows[hi]]
    data = [r for r in rows[hi + 1:] if any(_s(v) for v in r)]

    def idx(name):
        return hdr.index(name) if name in hdr else None

    i_code = idx('Κωδ. ΥΠΠΘ')
    # Το CSV του MySchool έχει 1-column shift μετά το col11 (βλ. MonadaDialog):
    # αν η στήλη «Κωδ. ΥΠΠΘ» δεν έχει κωδικούς ενώ η προηγούμενη έχει → shift.
    def _looks_code(col):
        vals = [_clean_code(r[col]) for r in data[:30] if col < len(r)]
        vals = [v for v in vals if v]
        return bool(vals) and sum(v.isdigit() and len(v) >= 6 for v in vals) / len(vals) > 0.8
    shift = 0
    if not _looks_code(i_code) and _looks_code(i_code - 1):
        shift = 1

    def col(name, shifted=True):
        i = idx(name)
        if i is None:
            return None
        return i - shift if (shifted and i >= 11) else i

    c_dimos, c_eidos = col('Δήμος', False), col('Είδος', False)
    c_code, c_name   = col('Κωδ. ΥΠΠΘ'), col('Ονομασία')
    c_stud, c_anast  = col('Ενεργοί Μαθητές'), col('Αναστολή')

    out = []
    for r in data:
        g = lambda c: r[c] if (c is not None and c < len(r)) else None
        eidos = _s(g(c_eidos))
        if 'Ιδιωτικ' in eidos or 'Ξένο' in eidos:
            continue
        if _s(g(c_anast)).upper() in ('TRUE', 'NAI', 'ΝΑΙ'):
            continue
        code = _clean_code(g(c_code))
        if not code:
            continue
        try:
            stud = int(float(_s(g(c_stud)) or 0))
        except ValueError:
            stud = None
        out.append({'code': code, 'name': _s(g(c_name)), 'dimos': _s(g(c_dimos)),
                    'eidos': eidos, 'students': stud})
    return pd.DataFrame(out).drop_duplicates('code')


def _yn(v):
    return _s(v).upper() in ('NAI', 'ΝΑΙ', 'YES', 'TRUE')


def _bkey(name):
    """Κλειδί σύγκρισης ονόματος κτιρίου (κενά/πεζά-κεφαλαία)."""
    return re.sub(r'\s+', ' ', _s(name)).casefold()


def load_24(path):
    """2.4 → dict {code: {bkey: {'name': κτίριο, E..K counts}}} μόνο για
    «Αίθουσα Διδασκαλίας», ανά κτίριο (στήλη «Κτίριο»)."""
    rows = _read_rows(path)
    hi = _find_header_row(rows, 'Κωδικός σχολείου')
    if hi is None:
        raise ValueError('Το 2.4 δεν έχει στήλη «Κωδικός σχολείου».')
    hdr = [_s(v) for v in rows[hi]]
    need = ['Κωδικός σχολείου', 'Χρήση αίθουσας', 'Διαθέτει ηλεκτρική εγκατάσταση',
            'Δίκτυο', 'WiFi', 'Ενεργός Χώρος', 'Έχει διαδραστικό πίνακα', 'Λυόμενο']
    miss = [n for n in need if n not in hdr]
    if miss:
        raise ValueError(f'Το 2.4 δεν έχει τις στήλες: {", ".join(miss)}')
    ix = {n: hdr.index(n) for n in need}
    i_bld = hdr.index('Κτίριο') if 'Κτίριο' in hdr else None
    res = {}
    for r in rows[hi + 1:]:
        g = lambda n: r[ix[n]] if ix[n] < len(r) else None
        if _s(g('Χρήση αίθουσας')) != 'Αίθουσα Διδασκαλίας':
            continue
        code = _clean_code(g('Κωδικός σχολείου'))
        if not code:
            continue
        bname = _s(r[i_bld]) if (i_bld is not None and i_bld < len(r)) else ''
        d = res.setdefault(code, {}).setdefault(
            _bkey(bname), {'name': bname, **dict.fromkeys('EFGHIJK', 0)})
        d['E'] += 1
        d['F'] += _yn(g('Ενεργός Χώρος'))
        d['G'] += _yn(g('Λυόμενο'))
        d['H'] += _yn(g('Διαθέτει ηλεκτρική εγκατάσταση'))
        d['I'] += _yn(g('Δίκτυο'))
        d['J'] += _yn(g('WiFi'))
        d['K'] += _yn(g('Έχει διαδραστικό πίνακα'))
    return res


def load_25(path):
    """2.5 → dict {code: [κτίριο, …]} με τη σειρά του αρχείου· κάθε κτίριο:
    {'name': Ονομασία Κτιρίου, 'L': έτος|None, 'M'..'AG': 'Ναι'/'Όχι'}."""
    rows = _read_rows(path)
    # Γραμμή υπο-κεφαλίδων: αυτή που έχει «Οικοδομικά Προβλήματα» στη θέση 6
    hi = None
    for i, r in enumerate(rows[:8]):
        if len(r) > 6 and _s(r[6]) == EXPECTED_25[6]:
            hi = i
    if hi is None:
        raise ValueError('Το 2.5 δεν έχει την αναμενόμενη μορφή '
                         '(στήλη «Οικοδομικά Προβλήματα» στη θέση G).')
    for pos, lbl in EXPECTED_25.items():
        if len(rows[hi]) <= pos or _s(rows[hi][pos]) != lbl:
            raise ValueError(f'Το 2.5: αναμενόταν «{lbl}» στη στήλη {get_column_letter(pos + 1)}.')
    out_cols = [get_column_letter(c) for c in range(13, 34)]   # M..AG
    res = {}
    for r in rows[hi + 1:]:
        if not r:
            continue
        code = _clean_code(r[0])
        if not code or ' ' in code:      # κενό ή κείμενο κεφαλίδας
            continue
        try:
            yr = int(float(_s(r[COL_25_YEAR]))) if _s(r[COL_25_YEAR]) else None
        except ValueError:
            yr = None
        d = {'name': _s(r[2]) if len(r) > 2 else '', 'L': yr}
        for pos, c in zip(COLS_25_YN, out_cols):
            v = _s(r[pos]) if pos < len(r) else ''
            d[c] = v if v in (YES, NO) else (v or None)
        res.setdefault(code, []).append(d)
    return res


# ═══════════════════════════════════════════════════════════════════
# ΕΠΕΞΕΡΓΑΣΙΑ
# ═══════════════════════════════════════════════════════════════════

_ORD_RE = re.compile(r'^\s*(\d+)\s*(?:ο|o|η)?\s*', re.IGNORECASE)


def _sort_key(row):
    """Νηπιαγωγεία → Δημοτικά· μέσα σε κάθε είδος: ανά περιοχή (όνομα χωρίς
    αριθμό) και μετά ανά αριθμό (1ο, 2ο, … 10ο)."""
    eidos_rank = 0 if 'Νηπιαγωγ' in row['eidos'] else 1
    name = row['name']
    m = _ORD_RE.match(name)
    num = int(m.group(1)) if m else 0
    rest = name[m.end():] if m else name
    return (eidos_rank, rest, num)


def process(path_22, path_24, path_25):
    """Επιστρέφει (df, warnings). df: ΜΙΑ ΓΡΑΜΜΗ ΑΝΑ ΚΤΙΡΙΟ με στήλες
    code, name, dimos, eidos, students, building, E..AG (στήλες προτύπου).
    Σχολείο χωρίς 2.5 → 1 γραμμή (κενά κτιριακά) με όλες τις αίθουσες."""
    df_sc = load_22(path_22)
    rooms = load_24(path_24)
    bld   = load_25(path_25)
    room_cols = list('EFGHIJK')
    bld_cols  = [get_column_letter(c) for c in range(12, 34)]   # L..AG
    df_sc['_key'] = df_sc.apply(_sort_key, axis=1)
    df_sc = df_sc.sort_values(['dimos', '_key']).drop(columns='_key')

    rows, no24, no25, orphan = [], [], [], []
    for _, sc in df_sc.iterrows():
        code = sc['code']
        blds = bld.get(code) or [None]
        r24  = rooms.get(code, {})
        if not r24:
            no24.append((code, sc['name']))
        if code not in bld:
            no25.append((code, sc['name']))
        # αίθουσες ανά κτίριο: ταίριασμα ονόματος· ό,τι περισσεύει → 1ο κτίριο
        keys = [_bkey(b['name']) if b else None for b in blds]
        per  = [dict.fromkeys(room_cols, 0) for _ in blds]
        for k, cnt in r24.items():
            tgt = keys.index(k) if k in keys else 0
            if k not in keys and len(blds) > 1:
                orphan.append((code, sc['name'],
                               f'Αίθουσες κτιρίου «{cnt["name"] or "(κενό)"}» του 2.4 δεν '
                               f'αντιστοιχούν σε κτίριο του 2.5 — μετρήθηκαν στο 1ο κτίριο'))
            for c in room_cols:
                per[tgt][c] += cnt[c]
        for b, pr in zip(blds, per):
            row = {'code': code, 'name': sc['name'], 'dimos': sc['dimos'],
                   'eidos': sc['eidos'], 'students': sc['students'],
                   'building': (b or {}).get('name', '')}
            for c in room_cols:
                row[c] = pr[c] if r24 else None
            for c in bld_cols:
                row[c] = (b or {}).get(c)
            rows.append(row)

    df = pd.DataFrame(rows, dtype=object)
    warnings = []
    if no24:
        warnings.append(f'{len(no24)} σχολεία χωρίς αίθουσες διδασκαλίας στο 2.4')
    if no25:
        warnings.append(f'{len(no25)} σχολεία χωρίς κτίριο στο 2.5')
    if orphan:
        warnings.append(f'{len(orphan)} περιπτώσεις αιθουσών με κτίριο που δεν υπάρχει στο 2.5')
    df.attrs['missing_24'] = no24
    df.attrs['missing_25'] = no25
    df.attrs['orphans'] = orphan
    return df, warnings


# ═══════════════════════════════════════════════════════════════════
# EXCEL
# ═══════════════════════════════════════════════════════════════════

def _school_year(d):
    y = d.year if d.month >= 9 else d.year - 1
    return f'{y}-{str(y + 1)[2:]}'


def _write_sheet(ws, df_s, title):
    """df_s: γραμμές ανά κτίριο (συνεχόμενες ανά σχολείο)."""
    from openpyxl.utils import column_index_from_string as _ci
    ncol = _ci(LAST_COL)
    F = lambda sz=11, b=False: Font(name='Calibri', size=sz, bold=b)

    # Τίτλος
    ws.merge_cells(f'A1:{LAST_COL}1')
    ws['A1'] = title
    ws['A1'].font = F(16, True)
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 51.75

    # Κεφαλίδες γραμμές 3–5
    for tpl, txt in [('A', 'Α/Α'), ('B', 'ΚΩΔ.'), ('C', 'ΣΧΟΛΕΙΟ'), ('D', 'ΜΑΘ. ΔΥΝΑΜΙΚΟ')]:
        col = oc(tpl)
        ws.merge_cells(f'{col}3:{col}5')
        ws[f'{col}3'] = txt
    ws.merge_cells(f'{BLD_COL}3:{BLD_COL}5')
    ws[f'{BLD_COL}3'] = 'ΚΤΙΡΙΟ'
    for col in ('A', 'B', 'C', BLD_COL, oc('D')):
        ws[f'{col}3'].font = F(11, True)
        ws[f'{col}3'].alignment = CTR
    ws.merge_cells(_shift_rng('E3:K4'))
    ws[f"{oc('E')}3"] = 'ΑΙΘΟΥΣΕΣ ΔΙΔΑΣΚΑΛΙΑΣ'
    ws.merge_cells(f"{oc('L')}3:{LAST_COL}3")
    ws[f"{oc('L')}3"] = 'ΚΤΙΡΙΑΚΑ'
    for col in (oc('E'), oc('L')):
        ws[f'{col}3'].font = F(14, True)
        ws[f'{col}3'].alignment = CTR
    for rng, txt, fill in HDR4:
        orng = _shift_rng(rng)
        a, b = orng.split(':')
        if a != b:
            ws.merge_cells(orng)
        ws[a].value = txt
        ws[a].font = F(11, True)
        ws[a].alignment = CTR
    for tpl, txt, fill in HDR5:
        c = ws[f'{oc(tpl)}5']
        c.value = txt
        c.font = F(11, True)
        c.alignment = VERT
        c.fill = fill
    e_i, k_i, l_i = _ci(oc('E')), _ci(oc('K')), _ci(oc('L'))
    for r in (3, 4, 5):
        for ci in range(1, ncol + 1):
            cell = ws.cell(r, ci)
            cell.border = BORDER
            if ci < e_i or (r == 3 and ci >= l_i):
                cell.fill = F_GREY
            elif e_i <= ci <= k_i:
                cell.fill = F_ORANGE
            elif r == 4:
                cl = get_column_letter(ci)
                cell.fill = next((f for rng, _, f in HDR4
                                  if _in_range(cl, _shift_rng(rng))), F_GREY)
    ws.row_dimensions[3].height = 18.75
    ws.row_dimensions[4].height = 35.25
    ws.row_dimensions[5].height = 162.75

    # Δεδομένα — μία γραμμή ανά κτίριο
    fill_out = {oc(c): f for c, f in FILL_BY_COL.items() if c not in ('L', 'M', 'N')}
    data_tpl = [get_column_letter(c) for c in range(5, 34)]          # E..AG
    first = 6
    r = first
    aa = 0
    for code, g in df_s.groupby('code', sort=False):
        aa += 1
        n = len(g)
        for j, (_, row) in enumerate(g.iterrows()):
            vals = {'A': aa, 'B': code, 'C': row['name'],
                    BLD_COL: row['building'] or None, oc('D'): row['students']}
            for t in data_tpl:
                vals[oc(t)] = row[t]
            for ci in range(1, ncol + 1):
                cl = get_column_letter(ci)
                cell = ws.cell(r, ci)
                if j == 0 or cl not in ('A', 'B', 'C', oc('D')):
                    cell.value = vals.get(cl)
                cell.font = F()
                cell.border = BORDER
                cell.alignment = Alignment(
                    horizontal='center' if (ci >= _ci(oc('M')) or cl in ('A', 'B')) else None,
                    vertical='center', wrap_text=(cl in ('C', BLD_COL)))
                if cl in fill_out:
                    cell.fill = fill_out[cl]
            ws.row_dimensions[r].height = 29.25
            r += 1
        if n > 1:
            for cl in ('A', 'B', 'C', oc('D')):
                ws.merge_cells(f'{cl}{r - n}:{cl}{r - 1}')
    last = r - 1

    # ΣΥΝΟΛΑ
    tr = last + 1
    ws.merge_cells(f'A{tr}:{BLD_COL}{tr}')
    ws[f'A{tr}'] = 'ΣΥΝΟΛΑ'
    for ci in range(1, ncol + 1):
        cl = get_column_letter(ci)
        cell = ws.cell(tr, ci)
        cell.font = F(14, True)
        cell.border = BORDER
        if cl in fill_out:
            cell.fill = fill_out[cl]
        if len(df_s) == 0 or ci <= _ci(BLD_COL):
            continue
        # ΣΥΝΟΛΑ μόνο σε ΜΑΘ. ΔΥΝΑΜΙΚΟ + ΑΙΘΟΥΣΕΣ· στα ΚΤΙΡΙΑΚΑ κενά
        # (άθροισμα/μέσος όρος δεν έχει νόημα — απόφαση χρήστη v4.5.0)
        if ci <= k_i:
            cell.value = f'=SUM({cl}{first}:{cl}{last})'
    ws[f'A{tr}'].alignment = Alignment(horizontal='right', vertical='center')
    ws.row_dimensions[tr].height = 18.75

    # Πλάτη, εκτύπωση
    for ci in range(1, 34):                      # στήλες προτύπου A..AG
        tpl = get_column_letter(ci)
        ws.column_dimensions[oc(tpl)].width = COL_WIDTHS.get(tpl, NARROW if ci >= 13 else 7)
    ws.column_dimensions[oc('L')].width = 8.3
    ws.column_dimensions[BLD_COL].width = BLD_WIDTH
    ws.freeze_panes = f"{oc('D')}6"
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '3:5'


def list_dimoi(df):
    """Ταξινομημένη λίστα Δήμων από το αποτέλεσμα του process()."""
    return sorted(d for d in df['dimos'].dropna().unique() if str(d).strip())


def build_for_dimos(df, dimos, out_path, ref_date=None):
    """Excel για ΕΝΑΝ Δήμο (1 φύλλο + ΕΛΛΕΙΨΕΙΣ αν χρειάζεται).
    dimos=None → όλοι οι Δήμοι (ΟΛΟΙ ΟΙ ΔΗΜΟΙ + φύλλο ανά Δήμο).
    Επιστρέφει (πλήθος σχολείων, πλήθος κτιρίων)."""
    if dimos is None:
        build_workbook(df, out_path, ref_date, all_sheet=True)
        return df['code'].nunique(), len(df)
    t = df[df['dimos'] == dimos].reset_index(drop=True)
    codes = set(t['code'])
    t.attrs = {k: [x for x in df.attrs.get(k, []) if x[0] in codes]
               for k in ('missing_24', 'missing_25', 'orphans')}
    build_workbook(t, out_path, ref_date, all_sheet=False)
    return t['code'].nunique(), len(t)


def _in_range(col_letter, rng):
    from openpyxl.utils import column_index_from_string as ci
    a, b = rng.split(':')
    a, b = re.sub(r'\d', '', a), re.sub(r'\d', '', b)
    return ci(a) <= ci(col_letter) <= ci(b)


def build_workbook(df, out_path, ref_date=None, all_sheet=True):
    """Ένα φύλλο ανά Δήμο (+ «ΟΛΟΙ ΟΙ ΔΗΜΟΙ» πρώτο αν all_sheet)."""
    ref_date = ref_date or datetime.today()
    sy = _school_year(ref_date)
    ds = ref_date.strftime('%d/%m/%Y')
    wb = Workbook()
    wb.remove(wb.active)
    if all_sheet:
        ws = wb.create_sheet('ΟΛΟΙ ΟΙ ΔΗΜΟΙ')
        _write_sheet(ws, df,
                     f'ΣΤΟΙΧΕΙΑ ΔΗΜΟΣΙΩΝ ΣΧ. ΜΟΝΑΔΩΝ \nΠ/ΘΜΙΑΣ ΕΚΠ/ΣΗΣ '
                     f'ΣΧ. ΕΤΟΥΣ {sy} (MYSCHOOL {ds})')
    for dimos, g in df.groupby('dimos', sort=True):
        name = re.sub(r'[\\/*?:\[\]]', '_', dimos)[:31] or 'ΧΩΡΙΣ ΔΗΜΟ'
        ws = wb.create_sheet(name)
        _write_sheet(ws, g.reset_index(drop=True),
                     f'ΣΤΟΙΧΕΙΑ ΔΗΜΟΣΙΩΝ ΣΧ. ΜΟΝΑΔΩΝ \nΠ/ΘΜΙΑΣ ΕΚΠ/ΣΗΣ ΔΗΜΟΥ '
                     f'{dimos} ΣΧ. ΕΤΟΥΣ {sy} (MYSCHOOL {ds})')
    # Φύλλο ελέγχου: σχολεία χωρίς στοιχεία 2.4/2.5
    miss = ([(c, n, 'Δεν βρέθηκαν αίθουσες διδασκαλίας στο 2.4') for c, n in df.attrs.get('missing_24', [])] +
            [(c, n, 'Δεν βρέθηκε κτίριο στο 2.5') for c, n in df.attrs.get('missing_25', [])] +
            list(df.attrs.get('orphans', [])))
    if miss:
        ws = wb.create_sheet('ΕΛΛΕΙΨΕΙΣ')
        for ci, h in enumerate(['ΚΩΔ.', 'ΣΧΟΛΕΙΟ', 'ΠΑΡΑΤΗΡΗΣΗ'], 1):
            c = ws.cell(1, ci, value=h)
            c.font = Font(bold=True)
            c.fill = F_GREY
            c.border = BORDER
        for ri, (c_, n, msg) in enumerate(miss, 2):
            for ci, v in enumerate([c_, n, msg], 1):
                ws.cell(ri, ci, value=v).border = BORDER
        ws.column_dimensions['A'].width = 10
        ws.column_dimensions['B'].width = 45
        ws.column_dimensions['C'].width = 80
    wb.save(out_path)


# ═══════════════════════════════════════════════════════════════════
# ΕΚΤΕΛΕΣΗ (CUSTOM_RUN — καλείται από το CheckRunDialog)
# ═══════════════════════════════════════════════════════════════════

def run(config):
    from core.framework import (get_downloaded_file, ask_file,
                                _missing_file_dialog, _show_results_popup)
    import core.framework as _fw
    _fw._current_check_title = CHECK_TITLE

    print('=' * 65)
    print(f'  {CHECK_TITLE}')
    print('=' * 65)
    # Πρώτα αυτόματα από τα σημερινά downloads· αν λείπει → επιλογή αρχείου
    # (π.χ. χειροκίνητη λήψη από το MySchool).
    paths = {}
    for rid, lbl in [('2.2', 'Εκτεταμένα Στοιχεία Σχολ. Μονάδων'),
                     ('2.4', 'Χώροι σχολικών μονάδων'),
                     ('2.5', 'Κτιριακά στοιχεία')]:
        p = get_downloaded_file(rid, f'Αρχείο {rid}:', silent=True)
        if not p:
            p = ask_file(f'Αρχείο {rid} — {lbl} [csv / xlsx / zip]:', required=False)
        paths[rid] = p
    path_22, path_24, path_25 = paths['2.2'], paths['2.4'], paths['2.5']
    if not (path_22 and path_24 and path_25):
        _missing_file_dialog(CHECK_TITLE, REQUIRED_REPORTS)
        return

    today = datetime.today()
    print('\nΕπεξεργασία...')
    df, warnings = process(path_22, path_24, path_25)
    print(f'  ✓ Σχολεία     : {df["code"].nunique()}  (κτίρια/γραμμές: {len(df)})')
    print(f'  ✓ Δήμοι       : {df["dimos"].nunique()}')
    for w in warnings:
        print(f'  ⚠ {w}')

    _docs   = os.path.join(os.path.expanduser('~'), 'Documents', 'MySchoolChecks')
    out_dir = os.path.join(_docs, f'results_{today.strftime("%Y%m%d")}', RESULTS_FOLDER)
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, f'{today.strftime("%Y%m%d")}_ΑΝΟΙΧΤΑ_ΔΕΔΟΜΕΝΑ.xlsx')
    build_workbook(df, out_path, today)
    print(f'  ✓ Αρχείο      : {os.path.basename(out_path)}')

    body = (f'Ημερομηνία: {today.strftime("%d/%m/%Y")}\n\n'
            f'Σχολεία: {df["code"].nunique()}  •  Κτίρια: {len(df)}  •  Δήμοι: {df["dimos"].nunique()}\n'
            f'(ένα φύλλο ανά Δήμο + «ΟΛΟΙ ΟΙ ΔΗΜΟΙ»)')
    if warnings:
        body += '\n\n⚠ ' + '\n⚠ '.join(warnings) + '\n(βλ. φύλλο «ΕΛΛΕΙΨΕΙΣ»)'
    _show_results_popup(CHECK_TITLE, body, result_type='warn' if warnings else 'ok',
                        excel_path=out_path)
